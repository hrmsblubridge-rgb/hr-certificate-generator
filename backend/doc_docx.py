"""Doc module — editable .docx output on the official letterhead.

Word cannot host a full-page PDF background, so the letterhead is split into
its header band and footer band (rasterised straight from the master PDF, never
redrawn) and placed in the section header / footer. Word then repeats them on
every page exactly like the PDF does. Page margins mirror the PDF safe area so
the text column lines up with the PDF output.
"""
from __future__ import annotations

import io
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from lxml import html as LH

from doc_builder import (
    DOC_DIR, LETTERHEAD, SIGNATURE, CONTENT_TOP, CONTENT_BOTTOM,
    CONTENT_LEFT, PAGE_W, PAGE_H, FOOTER_TOP, html_to_blocks,
    signature_table_html,
)

HEADER_BAND_PT = 96.0          # letterhead artwork above the header rule
HEADER_IMG = DOC_DIR / "_letterhead_header.png"
FOOTER_IMG = DOC_DIR / "_letterhead_footer.png"
BAND_DPI = 200

_ALIGN_MAP = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
}


def _ensure_bands():
    """Rasterise the letterhead's header / footer strips once and cache them."""
    if HEADER_IMG.is_file() and FOOTER_IMG.is_file():
        return
    with fitz.open(str(LETTERHEAD)) as doc:
        page = doc[0]
        zoom = BAND_DPI / 72.0
        matrix = fitz.Matrix(zoom, zoom)
        page.get_pixmap(matrix=matrix,
                        clip=fitz.Rect(0, 0, PAGE_W, HEADER_BAND_PT)).save(str(HEADER_IMG))
        page.get_pixmap(matrix=matrix,
                        clip=fitz.Rect(0, FOOTER_TOP - 6, PAGE_W, PAGE_H)).save(str(FOOTER_IMG))


def _add_runs(par, node, bold=False, italic=False, underline=False,
              sup=False, sub=False, strike=False):
    """Walk an inline HTML subtree and append equivalent Word runs."""
    def emit(text, **flags):
        if not text:
            return
        for i, chunk in enumerate(text.split("\n")):
            if i:
                par.add_run().add_break()
            if not chunk:
                continue
            run = par.add_run(chunk)
            run.bold = flags.get("bold")
            run.italic = flags.get("italic")
            run.underline = flags.get("underline")
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if flags.get("strike"):
                run.font.strike = True
            if flags.get("sup"):
                run.font.superscript = True
            if flags.get("sub"):
                run.font.subscript = True

    flags = dict(bold=bold, italic=italic, underline=underline,
                 sup=sup, sub=sub, strike=strike)
    emit(node.text, **flags)
    for child in node:
        tag = str(child.tag).lower() if isinstance(child.tag, str) else ""
        nxt = dict(flags)
        if tag in ("b", "strong"):
            nxt["bold"] = True
        elif tag in ("i", "em"):
            nxt["italic"] = True
        elif tag == "u":
            nxt["underline"] = True
        elif tag in ("s", "strike", "del"):
            nxt["strike"] = True
        elif tag == "sup":
            nxt["sup"] = True
        elif tag == "sub":
            nxt["sub"] = True
        elif tag == "br":
            par.add_run().add_break()
        if tag != "br":
            _add_runs(par, child, **nxt)
        emit(child.tail, **flags)


def _style_para(par, spacing_after=6):
    fmt = par.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(spacing_after)
    fmt.line_spacing = 1.15


def _add_html_paragraph(doc, el, tag: str):
    heading_sizes = {"h1": 15.5, "h2": 13.5, "h3": 12.5, "h4": 11.5}
    par = doc.add_paragraph()
    align = ""
    style_attr = (el.get("style") or "").lower()
    for candidate in ("center", "right", "justify", "left"):
        if f"text-align:{candidate}" in style_attr.replace(" ", ""):
            align = candidate
            break
    if align:
        par.alignment = _ALIGN_MAP[align]
    _add_runs(par, el, bold=tag in heading_sizes)
    if tag in heading_sizes:
        for run in par.runs:
            run.font.size = Pt(heading_sizes[tag])
            run.bold = True
    _style_para(par, spacing_after=6 if tag == "p" else 5)
    return par


def _add_table(doc, el):
    rows = el.xpath(".//tr")
    if not rows:
        return
    cols = max(len(r.xpath("./td|./th")) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    width = Inches(6.7 / cols)
    for tr in rows:
        cells = tr.xpath("./td|./th")
        row = table.add_row()
        for i in range(cols):
            cell = row.cells[i]
            cell.width = width
            cell._element.clear_content()
            if i >= len(cells):
                cell.add_paragraph()
                continue
            src = cells[i]
            inner_blocks = src.xpath("./p|./h1|./h2|./h3|./h4|./table")
            if inner_blocks:
                for blk in inner_blocks:
                    btag = str(blk.tag).lower()
                    if btag == "table":
                        continue
                    par = cell.add_paragraph()
                    _add_runs(par, blk, bold=str(src.tag).lower() == "th")
                    _style_para(par, spacing_after=3)
            else:
                par = cell.add_paragraph()
                _add_runs(par, src, bold=str(src.tag).lower() == "th")
                _style_para(par, spacing_after=3)


def build_docx(html_content: str, include_signature: bool = True,
               signature_style: str = "seal",
               signature_table: dict | None = None) -> bytes:
    """Editable Word version of the same document."""
    _ensure_bands()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.page_width = Pt(PAGE_W)
    section.page_height = Pt(PAGE_H)
    section.left_margin = Pt(CONTENT_LEFT)
    section.right_margin = Pt(CONTENT_LEFT)
    section.top_margin = Pt(CONTENT_TOP)
    section.bottom_margin = Pt(PAGE_H - CONTENT_BOTTOM)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(PAGE_H - FOOTER_TOP + 6)

    hdr = section.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hdr.paragraph_format.space_after = Pt(0)
    # Bleed the band to the physical page edges (headers respect the text
    # column by default, which would squeeze the artwork).
    hdr.paragraph_format.left_indent = Pt(-CONTENT_LEFT)
    hdr.paragraph_format.right_indent = Pt(-CONTENT_LEFT)
    hdr.add_run().add_picture(str(HEADER_IMG), width=Pt(PAGE_W))

    ftr = section.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ftr.paragraph_format.space_before = Pt(0)
    ftr.paragraph_format.left_indent = Pt(-CONTENT_LEFT)
    ftr.paragraph_format.right_indent = Pt(-CONTENT_LEFT)
    ftr.add_run().add_picture(str(FOOTER_IMG), width=Pt(PAGE_W))

    frag = LH.fragment_fromstring(
        "".join(b["html"] if b["kind"] != "table"
                else "<table>" + b["head"] + "".join(b["rows"]) + "</table>"
                for b in html_to_blocks(html_content)),
        create_parent="div")

    for el in frag:
        tag = str(el.tag).lower() if isinstance(el.tag, str) else ""
        if tag == "table":
            _add_table(doc, el)
            doc.add_paragraph()
        elif tag in ("ul", "ol"):
            for li in el.xpath("./li"):
                par = doc.add_paragraph(
                    style="List Bullet" if tag == "ul" else "List Number")
                _add_runs(par, li)
                _style_para(par, spacing_after=4)
        elif tag:
            _add_html_paragraph(doc, el, tag if tag.startswith("h") else "p")

    if include_signature:
        if signature_style == "table":
            data = signature_table or {}
            sig_html = signature_table_html(
                data.get("left") or {"label": "RECIPIENT / ADVISOR"},
                data.get("right") or {"label": "BLUBRIDGE TECHNOLOGIES PRIVATE LIMITED"},
                data.get("date") or "")
            _add_table(doc, LH.fragment_fromstring(sig_html))
        elif SIGNATURE.is_file():
            doc.add_paragraph()
            par = doc.add_paragraph()
            par.add_run().add_picture(str(SIGNATURE), width=Pt(200))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
